"""Milestone 5 — reports, exports, writeback, health (architecture §14, §18; prompt M5)."""

from __future__ import annotations

import csv
import io
import json
import os
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from waterflood_app.api.app import create_app, create_root_app
from waterflood_app.outputs import report as R
from waterflood_app.validation import synthetic_suite as suite

STREAK = suite.case_dir("streak_5x4")


@pytest.fixture(scope="module")
def env(tmp_path_factory: pytest.TempPathFactory) -> dict[str, object]:
    root = tmp_path_factory.mktemp("store")
    os.environ["WFO_ADMIN_PASSWORD"] = "admin-pass-123"
    client = TestClient(create_app(root, sync_jobs=True))

    def login(user: str, pw: str) -> dict[str, str]:
        r = client.post("/auth/login", json={"username": user, "password": pw})
        assert r.status_code == 200, r.text
        return {"Authorization": f"Bearer {r.json()['access_token']}"}

    admin = login("admin", "admin-pass-123")
    people: dict[str, dict[str, str]] = {"admin": admin}
    for key, roles in {
        "eng": ["engineer"],
        "rev": ["reviewer", "engineer"],
        "mgr": ["approver"],
        "ops": ["operations"],
        "view": ["viewer"],
    }.items():
        r = client.post(
            "/users",
            json={"username": f"{key}@asset.com", "password": f"pw-{key}", "roles": roles, "assets": ["ALPHA"]},
            headers=admin,
        )
        assert r.status_code == 201, r.text
        people[key] = login(f"{key}@asset.com", f"pw-{key}")
    r = client.post(
        "/projects", json={"name": "Streak pilot", "asset": "ALPHA", "unit_system": "field"}, headers=people["eng"]
    )
    pid = r.json()["id"]
    cid = client.post(
        "/connections", json={"project_id": pid, "kind": "files", "database": str(STREAK)}, headers=people["eng"]
    ).json()["id"]
    sug = client.get("/mapping/suggest", params={"connection_id": cid}, headers=people["eng"]).json()
    ds = {
        k: {"table": v["table"], "columns": v["columns"], "units": v["units"]}
        for k, v in sug["datasets"].items()
        if v["table"]
    }
    assert (
        client.post(
            "/mapping", json={"project_id": pid, "connection_id": cid, "datasets": ds}, headers=people["eng"]
        ).status_code
        == 200
    )
    job = client.post(
        "/runs",
        json={
            "project_id": pid,
            "variants": ["crmp"],
            "objective": "npv",
            "posture": "balanced",
            "economics": {"oil_price_usd_per_bbl": 65},
        },
        headers=people["rev"],
    ).json()
    st = client.get(f"/runs/jobs/{job['id']}", headers=people["rev"]).json()
    assert st["status"] == "done", st
    run_id = st["result_id"]
    rid = client.post("/recommendations", json={"run_id": run_id}, headers=people["eng"]).json()["id"]
    assert client.post(f"/recommendations/{rid}/review", json={"note": "ok"}, headers=people["rev"]).status_code == 200
    return {"client": client, "people": people, "pid": pid, "cid": cid, "run_id": run_id, "rid": rid, "root": root}


def test_reports_all_formats(env: dict[str, object]) -> None:
    client: TestClient = env["client"]  # type: ignore[assignment]
    people: dict[str, dict[str, str]] = env["people"]  # type: ignore[assignment]
    run_id = str(env["run_id"])
    html = client.get(f"/runs/{run_id}/report.html", headers=people["view"])
    assert html.status_code == 200 and html.headers["content-type"].startswith("text/html")
    body = html.text
    assert f"run {run_id}" in body  # watermark (§18)
    assert "<svg" in body and "Which method won" in body and "Change list" in body
    docx = client.get(f"/runs/{run_id}/report.docx", headers=people["view"])
    assert docx.status_code == 200 and docx.content[:2] == b"PK" and len(docx.content) > 50_000
    assert 'attachment; filename="wfo_run_' in docx.headers["content-disposition"]
    pdf = client.get(f"/runs/{run_id}/report.pdf", headers=people["view"])
    if R.pdf_renderer() is None:
        assert pdf.status_code == 409  # plain-language: no renderer on this host
    else:
        assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-", pdf.text[:200]
    # recommendation report carries the workflow section and the tornado
    rec_html = client.get(f"/recommendations/{env['rid']}/report.html", headers=people["view"]).text
    assert "Workflow" in rec_html and "REVIEWED" in rec_html and "Sensitivity of the gain" in rec_html
    assert client.get("/runs/nope/report.html", headers=people["view"]).status_code == 404
    assert client.get(f"/runs/{run_id}/report.txt", headers=people["view"]).status_code == 422
    # audit entries for reports
    audit = client.get("/admin/audit", params={"object_id": run_id}, headers=people["admin"]).json()
    assert any(a["action"] == "report" and a["detail"]["format"] == "html" for a in audit)


def test_exports(env: dict[str, object]) -> None:
    client: TestClient = env["client"]  # type: ignore[assignment]
    people: dict[str, dict[str, str]] = env["people"]  # type: ignore[assignment]
    run_id = str(env["run_id"])
    x = client.get(f"/runs/{run_id}/export.xlsx", headers=people["view"])
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(x.content), read_only=True)
    names = wb.sheetnames
    for sheet in (
        "README",
        "run",
        "wells",
        "change_list",
        "connectivity",
        "tau",
        "forecast_field",
        "history_match",
        "leaderboard",
        "gates",
        "conditions",
    ):
        assert sheet in names, names
    ws = wb["connectivity"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header[:4] == ["window", "sector", "injector", "producer"]
    readme = next(c.value for c in next(wb["README"].iter_rows(min_row=2, max_row=2)))
    assert f"run {run_id}" in str(readme)
    z = client.get(f"/runs/{run_id}/export.csv", headers=people["view"])
    assert z.status_code == 200 and z.headers["content-type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(z.content)) as zf:
        assert "change_list.csv" in zf.namelist() and "README.txt" in zf.namelist()
        rows = list(csv.DictReader(io.StringIO(zf.read("change_list.csv").decode("utf-8"))))
        assert rows and "rate_to" in rows[0]
    mj = client.get(f"/runs/{run_id}/export.json", headers=people["view"])
    assert mj.status_code == 200
    m = mj.json()
    assert m["run_id"] == run_id and m["sectors"][0]["variant"] == "crmp"
    assert len(m["sectors"][0]["f_ij"]) == len(m["sectors"][0]["injectors"])
    assert m["units"]["rate"] == "bbl/d"
    # recommendation exports add the workflow and rates sheets
    rx = client.get(f"/recommendations/{env['rid']}/export.xlsx", headers=people["view"])
    assert "rates" in load_workbook(io.BytesIO(rx.content), read_only=True).sheetnames
    # asset scope: another asset's user gets 403
    other = client.post(
        "/users",
        json={"username": "o@b.com", "password": "pw-o", "roles": ["viewer"], "assets": ["BRAVO"]},
        headers=people["admin"],
    )
    assert other.status_code == 201
    tok = client.post("/auth/login", json={"username": "o@b.com", "password": "pw-o"}).json()["access_token"]
    assert client.get(f"/runs/{run_id}/export.json", headers={"Authorization": f"Bearer {tok}"}).status_code == 403


def test_writeback_gate_and_targets(env: dict[str, object]) -> None:
    client: TestClient = env["client"]  # type: ignore[assignment]
    people: dict[str, dict[str, str]] = env["people"]  # type: ignore[assignment]
    root: Path = env["root"]  # type: ignore[assignment]
    rid, pid = str(env["rid"]), str(env["pid"])
    target_dir = root / "surveillance"
    wb_conn = client.post(
        "/connections",
        json={"project_id": pid, "kind": "files", "database": str(target_dir), "options": {"wfo_writeback": "1"}},
        headers=people["admin"],
    ).json()["id"]
    body = {"connection_id": wb_conn, "effective_date": "2026-10-01", "note": "October targets"}
    # not approved yet → 409; wrong role → 403
    assert client.post(f"/recommendations/{rid}/writeback", json=body, headers=people["mgr"]).status_code == 409
    assert client.post(f"/recommendations/{rid}/approve", json={"note": "go"}, headers=people["mgr"]).status_code == 200
    for who in ("eng", "rev", "ops", "view"):
        assert client.post(f"/recommendations/{rid}/writeback", json=body, headers=people[who]).status_code == 403, who
    # read connection (not flagged) → 409 with a plain message
    r = client.post(f"/recommendations/{rid}/writeback", json={"connection_id": env["cid"]}, headers=people["mgr"])
    assert r.status_code == 409 and "not enabled for writeback" in r.json()["detail"]
    r = client.post(f"/recommendations/{rid}/writeback", json=body, headers=people["mgr"])
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["kind"] == "csv" and out["n_rows"] == 5 and out["effective_date"] == "2026-10-01"
    csv_path = target_dir / "wfo_injection_targets.csv"
    rows = list(csv.DictReader(csv_path.open(encoding="utf-8")))
    assert len(rows) == 5 and rows[0]["recommendation_id"] == rid and rows[0]["approved_by"] == "mgr@asset.com"
    assert rows[0]["rate_unit"] == "bbl/d" and float(rows[0]["target_rate"]) > 0
    # second write appends; list shows both; audit + admin acting as approver works
    r2 = client.post(
        f"/recommendations/{rid}/writeback",
        json={**body, "table": "targets_v2"},
        headers={**people["admin"], "X-Acting-Role": "approver"},
    )
    assert r2.status_code == 201 and (target_dir / "targets_v2.csv").exists()
    lst = client.get(f"/recommendations/{rid}/writebacks", headers=people["view"]).json()
    assert [w["table"] for w in lst] == ["wfo_injection_targets", "targets_v2"]
    audit = client.get("/admin/audit", params={"object_id": rid}, headers=people["admin"]).json()
    assert any(a["action"] == "writeback" and a["acting_role"] == "approver" for a in audit)
    # SQL target (sqlite) creates the table and inserts rows
    db = root / "surveillance.sqlite"
    sql_conn = client.post(
        "/connections",
        json={"project_id": pid, "kind": "sqlite", "database": str(db), "options": {"wfo_writeback": "1"}},
        headers=people["admin"],
    ).json()["id"]
    r3 = client.post(f"/recommendations/{rid}/writeback", json={"connection_id": sql_conn}, headers=people["mgr"])
    assert r3.status_code == 201 and r3.json()["kind"] == "sql", r3.text
    import sqlite3

    n = sqlite3.connect(db).execute("SELECT COUNT(*) FROM wfo_injection_targets").fetchone()[0]
    assert n == 5
    # deployment switch
    svc = client.app.state.service  # type: ignore[attr-defined]
    svc.store.set_setting("writeback_enabled", False)
    assert client.post(f"/recommendations/{rid}/writeback", json=body, headers=people["mgr"]).status_code == 409
    svc.store.set_setting("writeback_enabled", True)


def test_health_and_readiness(env: dict[str, object]) -> None:
    client: TestClient = env["client"]  # type: ignore[assignment]
    h = client.get("/health").json()
    assert h["status"] == "ok" and h["version"] == "0.5.0"
    r = client.get("/health/ready")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "ready" and body["checks"]["store"]["ok"] and body["checks"]["filesystem"]["ok"]
    assert body["checks"]["jobs"]["mode"] == "sync" and body["checks"]["store"]["projects"] >= 1
    assert "pdf_renderer" in body["checks"]


def test_single_container_root_app(tmp_path: Path) -> None:
    """The offline profile serves the built UI at / and the API under /api from one process."""
    dist = tmp_path / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<html><body>wfo</body></html>", encoding="utf-8")
    (dist / "assets" / "app.js").write_text("console.log(1)", encoding="utf-8")
    os.environ["WFO_ADMIN_PASSWORD"] = "admin-pass-123"
    app = create_root_app(dist, root=tmp_path / "store", sync_jobs=True)
    c = TestClient(app)
    assert c.get("/api/health").json()["status"] == "ok"
    assert c.get("/").text.startswith("<html>")
    assert c.get("/workflow/abc").text.startswith("<html>")  # SPA fallback
    assert c.get("/assets/app.js").status_code == 200
    assert c.post("/api/auth/login", json={"username": "admin", "password": "admin-pass-123"}).status_code == 200
    with pytest.raises(FileNotFoundError):
        create_root_app(tmp_path / "missing", root=tmp_path / "store2", sync_jobs=True)


def test_figures_render_svg_and_png() -> None:
    from waterflood_app.outputs import figures as F

    fig = F.Fig(200, 120, title="t")
    ax = F.Axes(fig, 30, 10, 160, 90, (0, 10), (0, 100))
    ax.frame("x", "y", title="demo")
    fig.polyline([(ax.px(t), ax.py(t * 10)) for t in range(11)], F.OIL, 2, dash="3,2")
    fig.text(100, 110, "rotated", 9, rotate=-90)
    fig.circle(50, 50, 4, fill=F.WATER, opacity=0.5)
    svg = fig.to_svg()
    assert svg.startswith("<svg") and "polyline" in svg and 'stroke-dasharray="3,2"' in svg
    png = fig.to_png(1.5)
    assert png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) > 500
    assert F.nice_ticks(0, 1234) == [0.0, 250.0, 500.0, 750.0, 1000.0]
    assert F.fmt_num(1_260_000) == "1.3M" and F.fmt_num(12_500) == "12k" and F.fmt_num(0.123) == "0.12"
    t = F.tornado([{"label": "a", "low": 90, "high": 120}], 100, "USD")
    assert "a" in t.to_svg()
    empty = F.forecast_fan({"producers": [], "dates": []}, {"rate": "bbl/d"})
    assert "No forecast" in empty.to_svg()


def test_economic_tornado_shape(env: dict[str, object]) -> None:
    client: TestClient = env["client"]  # type: ignore[assignment]
    people: dict[str, dict[str, str]] = env["people"]  # type: ignore[assignment]
    rec = client.get(f"/recommendations/{env['rid']}", headers=people["view"]).json()
    tor = rec["recommendation"]["tornado"]
    assert tor["unit"] == "USD" and len(tor["items"]) == 5
    labels = [i["label"] for i in tor["items"]]
    assert any("Oil price" in x for x in labels) and any("Forecast range" in x for x in labels)
    price = next(i for i in tor["items"] if "Oil price" in i["label"])
    assert price["low"] != price["high"]
    # expected oil gain in the action list is a volume even for the NPV objective
    acts = rec["recommendation"]["actions"]
    assert acts and all(a["expected_oil_gain"] >= 0 for a in acts) and sum(a["expected_oil_gain"] for a in acts) > 0
    assert json.dumps(tor)  # serialisable
