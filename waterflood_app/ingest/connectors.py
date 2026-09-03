"""File connectors and the load pipeline — architecture §5 (L1).

``read_tables`` turns a folder of CSV/Parquet files or one Excel workbook into named tables;
``load_project`` runs table suggestion → column mapping → apply → ID checks and returns a
:class:`LoadedData` with the canonical frames and every Condition raised on the way.
SQL connectors (read-only credentials in the backend) arrive with the API milestone.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import polars as pl

from waterflood_app.ingest import schema
from waterflood_app.ingest.mapping import (
    apply_mapping,
    report_unmatched_ids,
    suggest_mapping,
    suggest_table,
)
from waterflood_app.messaging.conditions import ConditionLog


def read_tables(source: Path) -> dict[str, pl.DataFrame]:
    """A folder → one table per CSV/Parquet file (stem = table name); a workbook → one per sheet."""
    source = Path(source)
    tables: dict[str, pl.DataFrame] = {}
    if source.is_dir():
        for p in sorted(source.iterdir()):
            if p.suffix.lower() == ".csv":
                tables[p.stem] = pl.read_csv(p, infer_schema_length=10000, try_parse_dates=False)
            elif p.suffix.lower() == ".parquet":
                tables[p.stem] = pl.read_parquet(p)
        return tables
    if source.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
        for name in wb.sheetnames:
            rows = list(wb[name].iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(h) if h is not None else f"col{k}" for k, h in enumerate(rows[0])]
            data = [list(r) for r in rows[1:] if any(v is not None for v in r)]
            tables[name] = pl.DataFrame(data, schema=header, orient="row", strict=False)
        return tables
    if source.suffix.lower() == ".csv":
        tables[source.stem] = pl.read_csv(source, infer_schema_length=10000)
        return tables
    raise ValueError(f"unsupported source {source}")


@dataclass
class LoadedData:
    rates: pl.DataFrame
    pressure: pl.DataFrame | None = None
    coords: pl.DataFrame | None = None
    category: pl.DataFrame | None = None
    events: pl.DataFrame | None = None
    mappings: dict[str, schema.ColumnMapping] = field(default_factory=dict)
    units: dict[str, dict[str, str]] = field(default_factory=dict)
    conditions: ConditionLog = field(default_factory=ConditionLog)
    source: str = ""

    def frame(self, name: str) -> pl.DataFrame | None:
        return {
            "rates": self.rates,
            "pressure": self.pressure,
            "coords": self.coords,
            "category": self.category,
            "events": self.events,
        }[name]


def load_project(
    source: Path,
    mappings: dict[str, schema.ColumnMapping] | None = None,
    unit_system: str = "field",
    aliases: dict[str, str] | None = None,
    tables: dict[str, pl.DataFrame] | None = None,
) -> LoadedData:
    """Load, map and validate the canonical datasets from files (or pre-read ``tables``)."""
    tables = tables if tables is not None else read_tables(source)
    log = ConditionLog()
    mappings = dict(mappings or {})
    frames: dict[str, pl.DataFrame | None] = {}
    units: dict[str, dict[str, str]] = {}
    names = list(tables)
    for ds in schema.DATASETS.values():
        m = mappings.get(ds.name)
        if m is None:
            table = suggest_table(ds, names)
            if table is None and ds.name in tables:
                table = ds.name
            if table is None:
                if ds.required:
                    raise ValueError(f"no table found for the required dataset {ds.name!r} in {names}")
                frames[ds.name] = None
                continue
            m = suggest_mapping(ds, tables[table].columns, unit_system)
            m.table = table
            mappings[ds.name] = m
        if m.table is None or m.table not in tables:
            frames[ds.name] = None
            continue
        applied = apply_mapping(ds, tables[m.table], m, log, aliases)
        frames[ds.name] = applied.frame
        units[ds.name] = applied.units
    rates = frames["rates"]
    assert rates is not None
    for name in ("pressure", "coords", "category", "events"):
        f = frames.get(name)
        if f is not None:
            report_unmatched_ids(rates, f, name, log)
    return LoadedData(
        rates=rates,
        pressure=frames.get("pressure"),
        coords=frames.get("coords"),
        category=frames.get("category"),
        events=frames.get("events"),
        mappings=mappings,
        units=units,
        conditions=log,
        source=str(source),
    )


def mapping_summary(loaded: LoadedData) -> dict[str, Any]:
    """Plain summary for the Details panel / tests: table, columns, units, per dataset."""
    return {
        ds: {"table": m.table, "columns": dict(m.columns), "units": dict(m.units)} for ds, m in loaded.mappings.items()
    }
