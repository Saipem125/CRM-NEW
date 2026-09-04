"""Tabular exports — architecture §14 "tables (XLSX/CSV), model JSON"; §18 exports watermarked with run ID.

All tables are built from the stored result bundle (display units) and the registry row, so
exports are available for any run without the in-memory engine objects.
"""

from __future__ import annotations

import io
import json
import zipfile
from datetime import UTC, datetime
from typing import Any

import polars as pl

from waterflood_app.outputs.report import APP_VERSION


def _float(v: Any) -> float | None:
    try:
        return None if v is None else float(v)
    except (TypeError, ValueError):
        return None


def tables_from_bundle(
    run_row: dict[str, Any], bundle: dict[str, Any], recommendation: dict[str, Any] | None = None
) -> dict[str, pl.DataFrame]:
    """Every table the XLSX/CSV export carries, keyed by sheet name."""
    units = dict(bundle.get("units") or {})
    rate_u, vol_u = units.get("rate", "bbl/d"), units.get("volume", "bbl")
    types = {w: t["type"] for w, t in (bundle.get("well_types") or {}).items()}
    latest = int(bundle.get("latest_window_index", 0))
    out: dict[str, pl.DataFrame] = {}
    summary = run_row.get("summary") or {}
    req = bundle.get("request") or summary.get("request") or {}
    out["run"] = pl.DataFrame(
        {
            "item": [
                "run_id",
                "project_id",
                "data_hash",
                "config_hash",
                "code_version",
                "seed",
                "created_at",
                "confidence",
                "objective",
                "posture",
                "horizon_months",
                "rate_unit",
                "volume_unit",
                "pressure_unit",
                "exported_at",
                "application",
            ],
            "value": [
                str(run_row.get("id", "")),
                str(run_row.get("project_id", bundle.get("project_id", ""))),
                str(bundle.get("data_hash", "")),
                str(bundle.get("config_hash", "")),
                str(run_row.get("code_version", "")),
                str(bundle.get("seed", "")),
                str(run_row.get("created_at", "")),
                str(bundle.get("confidence", "")),
                str(req.get("objective", "oil")),
                str(req.get("posture") or ""),
                str(req.get("horizon_months") or ""),
                rate_u,
                vol_u,
                units.get("pressure", "psi"),
                datetime.now(UTC).isoformat(timespec="seconds"),
                f"Waterflood Optimizer {APP_VERSION}",
            ],
        }
    )
    dq = bundle.get("data_quality") or {}
    cov = dq.get("coverage") or {}
    pp = (bundle.get("pressure") or {}).get("per_producer", {})
    sector_of: dict[str, str] = {}
    xy: dict[str, list[float | None]] = {}
    for s in bundle.get("sectors", []):
        if int(s.get("window_index", 0)) != latest:
            continue
        for w in s["injectors"] + s["producers"]:
            sector_of[w["well"]] = str(s["id"])
            xy[w["well"]] = list(w["xy"]) if w.get("xy") else [None, None]
    out["wells"] = pl.DataFrame(
        {
            "well": list(cov),
            "type": [types.get(w, "") for w in cov],
            "sector": [sector_of.get(w, "") for w in cov],
            "x": [xy.get(w, [None, None])[0] for w in cov],
            "y": [xy.get(w, [None, None])[1] for w in cov],
            "coverage_fraction": [float(v) for v in cov.values()],
            "pressure_source": [pp.get(w, "") for w in cov],
        }
    )
    change_rows: list[dict[str, Any]] = []
    conn_rows: list[dict[str, Any]] = []
    tau_rows: list[dict[str, Any]] = []
    eff_rows: list[dict[str, Any]] = []
    ff_rows: list[dict[str, Any]] = []
    fw_rows: list[dict[str, Any]] = []
    hm_rows: list[dict[str, Any]] = []
    lb_rows: list[dict[str, Any]] = []
    gate_rows: list[dict[str, Any]] = []
    for s in bundle.get("sectors", []):
        key, sid = str(s["key"]), str(s["id"])
        m = s.get("model") or {}
        inj = [w["well"] for w in s["injectors"]]
        prod = [w["well"] for w in s["producers"]]
        F_ = m.get("f_ij") or []
        C_ = m.get("pair_confidence") or []
        D_ = s.get("distances") or []
        for i, wi in enumerate(inj):
            for j, wp in enumerate(prod):
                conn_rows.append(
                    {
                        "window": key,
                        "sector": sid,
                        "injector": wi,
                        "producer": wp,
                        "f_ij": _float(F_[i][j]) if F_ and i < len(F_) and j < len(F_[i]) else None,
                        "pair_confidence": _float(C_[i][j]) if C_ and i < len(C_) and j < len(C_[i]) else None,
                        "distance_m": _float(D_[i][j]) if D_ and i < len(D_) and j < len(D_[i]) else None,
                    }
                )
        taus = m.get("tau_days") or []
        J = m.get("J")
        for j, wp in enumerate(prod):
            jv = J[j] if isinstance(J, list) and j < len(J) and not isinstance(J[j], list) else None
            tau_rows.append(
                {
                    "window": key,
                    "sector": sid,
                    "producer": wp,
                    "tau_days": _float(taus[j]) if j < len(taus) else None,
                    "J": _float(jv),
                    "blind_r2": _float(s["producers"][j].get("blind_r2")),
                    "blind_mape_pct": _float(s["producers"][j].get("blind_mape")),
                }
            )
        for r in s.get("leaderboard") or []:
            lb_rows.append(
                {
                    "window": key,
                    "sector": sid,
                    **{k: (v if not isinstance(v, list | dict) else json.dumps(v)) for k, v in r.items()},
                }
            )
        for g, ok in (s.get("gates") or {}).items():
            gate_rows.append({"window": key, "sector": sid, "gate": g, "pass": bool(ok)})
        dates = list(s.get("dates") or [])
        for p in s["producers"]:
            for t, d in enumerate(dates):
                hm_rows.append(
                    {
                        "sector": sid,
                        "window": key,
                        "date": d,
                        "producer": p["well"],
                        f"raw_liquid_{rate_u}": _float(p["raw_liquid"][t]),
                        f"cleaned_liquid_{rate_u}": _float(p["liquid"][t]),
                        f"model_liquid_{rate_u}": _float(p["model"][t]),
                        f"oil_{rate_u}": _float(p["oil"][t]),
                        f"water_{rate_u}": _float(p["water"][t]),
                        "bhp": _float(p["bhp"][t]) if p.get("bhp") else None,
                        "blind": t >= int(s.get("blind_start_index", len(dates))),
                    }
                )
        if int(s.get("window_index", 0)) != latest:
            continue
        rec = s.get("recommendation") or {}
        if recommendation is not None and recommendation.get("sector") == sid and recommendation.get("recommendation"):
            rec = recommendation["recommendation"]
        for a in rec.get("actions") or []:
            change_rows.append({"sector": sid, **{k: a.get(k) for k in a}})
        fc = s.get("forecast") or {}
        eff = fc.get("injector_efficiency")
        if eff:
            for k, w in enumerate(eff["injectors"]):
                eff_rows.append(
                    {
                        "sector": sid,
                        "injector": w,
                        f"oil_per_{vol_u}_injected": _float(eff["oil_per_bbl"][k]),
                        f"current_rate_{rate_u}": _float(eff["current"][k]),
                        f"optimized_rate_{rate_u}": _float(eff["optimized"][k]),
                        f"current_oil_{rate_u}": _float(eff["current_oil"][k]),
                        f"optimized_oil_{rate_u}": _float(eff["optimized_oil"][k]),
                    }
                )
        if fc:
            fp, fb = fc["field_plan"], fc["field_base"]
            for t, d in enumerate(fc["dates"]):
                ff_rows.append(
                    {
                        "sector": sid,
                        "date": d,
                        f"plan_p10_{rate_u}": _float(fp["p10"][t]),
                        f"plan_p50_{rate_u}": _float(fp["p50"][t]),
                        f"plan_p90_{rate_u}": _float(fp["p90"][t]),
                        f"base_p10_{rate_u}": _float(fb["p10"][t]),
                        f"base_p50_{rate_u}": _float(fb["p50"][t]),
                        f"base_p90_{rate_u}": _float(fb["p90"][t]),
                    }
                )
            plan, base = fc["plan"], fc["base"]
            for j, wp in enumerate(prod):
                for t, d in enumerate(fc["dates"]):
                    fw_rows.append(
                        {
                            "sector": sid,
                            "producer": wp,
                            "date": d,
                            f"plan_p50_{rate_u}": _float(plan["p50"][t][j]),
                            f"plan_p10_{rate_u}": _float(plan["p10"][t][j]),
                            f"plan_p90_{rate_u}": _float(plan["p90"][t][j]),
                            f"base_p50_{rate_u}": _float(base["p50"][t][j]),
                        }
                    )
    out["change_list"] = _frame(change_rows)
    out["connectivity"] = _frame(conn_rows)
    out["tau"] = _frame(tau_rows)
    out["injector_efficiency"] = _frame(eff_rows)
    out["forecast_field"] = _frame(ff_rows)
    out["forecast_wells"] = _frame(fw_rows)
    out["history_match"] = _frame(hm_rows)
    out["leaderboard"] = _frame(lb_rows)
    out["gates"] = _frame(gate_rows)
    out["conditions"] = _frame(
        [
            {
                "scope": c.get("scope"),
                "code": c.get("code"),
                "severity": c.get("severity"),
                "message": c.get("message"),
                "action": c.get("action"),
                "technical": c.get("technical"),
            }
            for c in bundle.get("conditions", [])
        ]
    )
    out["data_quality"] = pl.DataFrame(
        {"item": [k for k in dq if k != "coverage"], "value": [str(dq[k]) for k in dq if k != "coverage"]}
    )
    if recommendation is not None:
        out["workflow"] = _frame(
            [
                {
                    "at": h.get("at"),
                    "action": h.get("action"),
                    "actor": h.get("actor"),
                    "acting_role": h.get("acting_role"),
                    "from_state": h.get("from_state"),
                    "to_state": h.get("to_state"),
                    "note": h.get("note"),
                }
                for h in recommendation.get("history", [])
            ]
        )
        out["rates"] = _frame(
            [
                {
                    "injector": w,
                    f"recommended_{rate_u}": _float(v),
                    f"implemented_{rate_u}": _float((recommendation.get("implemented_rates") or {}).get(w)),
                }
                for w, v in (recommendation.get("recommended_rates") or {}).items()
            ]
        )
    return out


def _frame(rows: list[dict[str, Any]]) -> pl.DataFrame:
    if not rows:
        return pl.DataFrame({"empty": ["(no rows)"]})
    return pl.DataFrame(rows, infer_schema_length=None)


def to_xlsx(tables: dict[str, pl.DataFrame], watermark: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "README"
    ws0["A1"] = "Waterflood Optimizer export"
    ws0["A1"].font = Font(bold=True, size=13)
    ws0["A2"] = watermark
    ws0["A3"] = "Sheets: " + ", ".join(tables)
    ws0["A4"] = "Column suffixes carry the unit (e.g. _bbl/d); dates are ISO; f_ij are dimensionless; tau in days."
    ws0.column_dimensions["A"].width = 110
    head_fill = PatternFill("solid", fgColor="F3F5F8")
    for name, df in tables.items():
        ws = wb.create_sheet(title=name[:31])
        cols = df.columns
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            ws.cell(row=1, column=c).font = Font(bold=True)
            ws.cell(row=1, column=c).fill = head_fill
        for row in df.iter_rows():
            ws.append([_cell(v) for v in row])
        ws.freeze_panes = "A2"
        for k, col in enumerate(cols, start=1):
            width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200).to_list())) if df.height else len(col)
            ws.column_dimensions[get_column_letter(k)].width = min(max(10, width + 2), 60)
        ws.oddFooter.center.text = watermark
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v: Any) -> Any:
    if isinstance(v, list | dict):
        return json.dumps(v)
    return v


def to_csv_zip(tables: dict[str, pl.DataFrame], watermark: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(
            "README.txt",
            f"Waterflood Optimizer export\n{watermark}\n\nFiles: " + ", ".join(f"{n}.csv" for n in tables) + "\n",
        )
        for name, df in tables.items():
            z.writestr(f"{name}.csv", df.write_csv())
    return buf.getvalue()


def model_json(run_row: dict[str, Any], bundle: dict[str, Any]) -> dict[str, Any]:
    """Fitted model parameters per sector and window — the reproducible artefact of a run (§5, §14)."""
    units = dict(bundle.get("units") or {})
    sectors = []
    for s in bundle.get("sectors", []):
        m = s.get("model") or {}
        sectors.append(
            {
                "window": s["key"],
                "window_index": s.get("window_index"),
                "sector": s["id"],
                "dates": [s["dates"][0], s["dates"][-1]] if s.get("dates") else None,
                "blind_start_index": s.get("blind_start_index"),
                "injectors": [w["well"] for w in s["injectors"]],
                "producers": [w["well"] for w in s["producers"]],
                "variant": m.get("variant"),
                "label": m.get("label"),
                "f_ij": m.get("f_ij"),
                "tau_days": m.get("tau_days"),
                "tau_ij_days": m.get("tau_ij_days"),
                "J": m.get("J"),
                "J_unit": f"{units.get('rate', 'bbl/d')} per {units.get('pressure', 'psi')}",
                "extra": m.get("extra"),
                "n_params": m.get("n_params"),
                "metrics": {
                    "blind_r2": m.get("blind_r2"),
                    "blind_mape": m.get("blind_mape"),
                    "aicc": m.get("aicc"),
                    "spread": m.get("spread"),
                    "plausible": m.get("plausible"),
                },
                "confidence": s.get("confidence"),
                "gates": s.get("gates"),
                "leaderboard": s.get("leaderboard"),
            }
        )
    return {
        "application": f"Waterflood Optimizer {APP_VERSION}",
        "run_id": run_row.get("id"),
        "project_id": run_row.get("project_id", bundle.get("project_id")),
        "created_at": run_row.get("created_at"),
        "data_hash": bundle.get("data_hash"),
        "config_hash": bundle.get("config_hash"),
        "code_version": run_row.get("code_version"),
        "seed": bundle.get("seed"),
        "units": bundle.get("units"),
        "confidence": bundle.get("confidence"),
        "registry_models": run_row.get("models", []),  # internal-unit params as stored by the registry
        "sectors": sectors,
    }
